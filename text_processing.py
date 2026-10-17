"""Readable-content extraction and API-limit-safe narration planning."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import ipaddress
import re
import socket
from urllib.parse import urlparse, urljoin

import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

MAX_INPUT_CHARS = 80000
MAX_FILE_BYTES = 10 * 1024 * 1024


@dataclass(frozen=True)
class Speech:
    text: str


@dataclass(frozen=True)
class Pause:
    seconds: float


Event = Speech | Pause


def normalize_text(text: str) -> str:
    """Preserve Indic scripts and paragraph breaks, remove invisible noise."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[\u200b\ufeff\u2060]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def markdown_to_text(raw: str) -> str:
    raw = re.sub(r"```[\s\S]*?```", "", raw)
    raw = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", raw)
    raw = re.sub(r"\[([^\]]+)\]\((?:[^()]|\([^()]*\))*\)", r"\1", raw)
    raw = re.sub(r"^\s{0,3}#{1,6}\s+", "", raw, flags=re.M)
    raw = re.sub(r"^\s{0,3}>\s?", "", raw, flags=re.M)
    raw = re.sub(r"^\s*[-*+]\s+", "", raw, flags=re.M)
    raw = re.sub(r"(?<!\w)[*_]{1,2}(?=\S)|(?<=\S)[*_]{1,2}(?!\w)", "", raw)
    return normalize_text(raw)


def html_to_text(raw: str) -> str:
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg", "iframe"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    blocks = []
    for tag in root.find_all(["h1", "h2", "h3", "p", "li", "blockquote"]):
        # Avoid duplicate text inside nested blocks.
        if tag.find_parent(["p", "li", "blockquote"]) is not None:
            continue
        txt = tag.get_text(" ", strip=True)
        if txt and len(txt) > (5 if tag.name in ("p", "li") else 0):
            blocks.append(txt)
    if not blocks:
        return normalize_text(root.get_text("\n", strip=True))
    return normalize_text("\n\n".join(blocks))


def extract_file(filename: str, contents: bytes) -> str:
    """Extract native/selectable text; scanned PDFs need OCR outside this app."""
    if len(contents) > MAX_FILE_BYTES:
        raise ValueError("Document is larger than 10 MB.")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        reader = PdfReader(BytesIO(contents))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs aren't supported.")
        if len(reader.pages) > 200:
            raise ValueError("PDF has more than 200 pages; please split it first.")
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        if len(text.strip()) < 30:
            raise ValueError("No selectable text found. This may be a scanned PDF and needs OCR.")
    elif suffix == "docx":
        doc = Document(BytesIO(contents))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also include useful text in tables.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(". ".join(cells))
        text = "\n\n".join(paras)
    elif suffix in {"txt", "md", "markdown", "html", "htm"}:
        try:
            text = contents.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = contents.decode("cp1252", errors="replace")
        if suffix in {"html", "htm"}:
            text = html_to_text(text)
        elif suffix in {"md", "markdown"}:
            text = markdown_to_text(text)
    else:
        raise ValueError(f"Unsupported document type: .{suffix}")
    text = normalize_text(text)
    if not text:
        raise ValueError("No readable text found in this file.")
    return text


def validate_public_url(url: str) -> str:
    """Basic SSRF protection for locally run URL fetching (not a full hosted-service guard)."""
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"https", "http"} or not parsed.hostname:
        raise ValueError("Enter a complete http:// or https:// article URL.")
    if parsed.username or parsed.password:
        raise ValueError("URLs with embedded credentials aren't accepted.")
    if parsed.port and parsed.port not in {80, 443}:
        raise ValueError("Only standard web ports (80 and 443) are supported.")
    host = parsed.hostname.rstrip(".").lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
        raise ValueError("Local-network URLs aren't allowed.")
    try:
        addrinfos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError("Could not resolve this website.") from exc
    if not addrinfos:
        raise ValueError("Could not resolve this website.")
    for result in addrinfos:
        addr = ipaddress.ip_address(result[4][0])
        if not addr.is_global:
            raise ValueError("Private or reserved network URLs aren't allowed.")
    return url.strip()


def extract_url(url: str) -> str:
    url = validate_public_url(url)
    for _ in range(4):
        resp = requests.get(url, timeout=(8, 20), allow_redirects=False,
                            stream=True, headers={"User-Agent": "NarrateAIReader/1.0 (+text extraction)"})
        try:
            if resp.is_redirect:
                target = resp.headers.get("Location")
                if not target:
                    raise ValueError("Article redirect had no destination.")
                url = validate_public_url(urljoin(url, target))
                continue
            resp.raise_for_status()
            content_type = resp.headers.get("Content-Type", "").lower()
            if "text/html" not in content_type and "application/xhtml" not in content_type:
                raise ValueError("This URL doesn't serve an HTML article. Upload the file instead.")
            data = bytearray()
            for chunk in resp.iter_content(chunk_size=65536):
                data.extend(chunk)
                if len(data) > 2_000_000:
                    raise ValueError("Article page is too large (over 2 MB).")
            resp.encoding = resp.encoding or "utf-8"
            extracted = html_to_text(bytes(data).decode(resp.encoding, errors="replace"))
            if len(extracted) < 60:
                raise ValueError("Could not extract enough text. Try pasting the article instead.")
            return extracted
        finally:
            resp.close()
    raise ValueError("Too many redirects.")


def _split_long(text: str, limit: int) -> list[str]:
    """Split at punctuation, then whitespace, then characters as a last resort."""
    result: list[str] = []
    sentences = re.split(r"(?<=[.!?।॥])\s+", text.strip())
    current = ""
    def append_piece(piece: str) -> None:
        nonlocal current
        if not piece:
            return
        candidate = (current + " " + piece).strip() if current else piece
        if len(candidate) <= limit:
            current = candidate
        else:
            if current:
                result.append(current)
            current = piece

    for sentence in sentences:
        if len(sentence) <= limit:
            append_piece(sentence)
            continue
        for word in sentence.split():
            if len(word) > limit:
                if current:
                    result.append(current)
                    current = ""
                for i in range(0, len(word), limit):
                    result.append(word[i:i + limit])
            else:
                append_piece(word)
    if current:
        result.append(current)
    return result


def plan_narration(raw: str, limit: int = 2400, paragraph_pause: float = 0.45) -> list[Event]:
    """Plan batches with explicit [pause=0.8] and [em]word[/em] cues.

    [em] is a punctuation cue, NOT guaranteed SSML/prosody control.
    """
    if not 100 <= limit <= 2500:
        raise ValueError("Chunk limit must be between 100 and 2500 characters.")
    if not 0 <= paragraph_pause <= 5:
        raise ValueError("Paragraph pause must be between 0 and 5 seconds.")
    raw = normalize_text(raw)
    if not raw:
        return []
    if len(raw) > MAX_INPUT_CHARS:
        raise ValueError(f"Content exceeds {MAX_INPUT_CHARS:,} characters. Split it into smaller projects.")
    # Explicit phrase emphasis via punctuation. Avoid duplicated exclamation marks.
    raw = re.sub(r"\[em\]([\s\S]*?)\[/em\][.!?]?", lambda m: (m.group(1).strip().rstrip(".!?") + "!"), raw, flags=re.I)
    markers = re.compile(r"\[pause(?:\s*=\s*(\d+(?:\.\d+)?))?\]|\n\s*\n+", re.I)
    events: list[Event] = []
    previous_end = 0
    for match in markers.finditer(raw):
        text = re.sub(r"\s+", " ", raw[previous_end:match.start()]).strip()
        events.extend(Speech(chunk) for chunk in _split_long(text, limit) if chunk)
        if match.group().lower().startswith("[pause"):
            sec = min(float(match.group(1) or 0.7), 5.0)
        else:
            sec = paragraph_pause
        if sec and events and isinstance(events[-1], Speech):
            events.append(Pause(sec))
        previous_end = match.end()
    tail = re.sub(r"\s+", " ", raw[previous_end:]).strip()
    events.extend(Speech(chunk) for chunk in _split_long(tail, limit) if chunk)
    while events and isinstance(events[-1], Pause):
        events.pop()
    return events
